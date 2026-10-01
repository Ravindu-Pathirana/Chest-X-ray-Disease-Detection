"""Generate the T45 presentation from committed, identified source artifacts.

Install requirements-presentation.txt first. No model inference occurs here.
"""
from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/presentation/T45_research_progress.pptx"

NAVY = RGBColor(14, 26, 45)
PANEL = RGBColor(24, 43, 65)
CYAN = RGBColor(80, 215, 216)
WHITE = RGBColor(245, 248, 250)
MUTED = RGBColor(174, 190, 204)
CORAL = RGBColor(255, 155, 117)


def text_box(slide, x, y, w, h, text, *, size=20, color=WHITE, bold=False,
             align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = "Aptos"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return shape


def rect(slide, x, y, w, h, fill=PANEL, radius=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def base(prs, title, kicker="RESEARCH PROGRESS"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    rect(slide, 0, 0, 0.13, 7.5, CYAN)
    text_box(slide, 0.55, 0.28, 7, 0.35, kicker, size=11, color=CYAN, bold=True)
    text_box(slide, 0.55, 0.76, 12, 0.85, title, size=31, bold=True)
    rect(slide, 0.55, 1.67, 12.2, 0.025, CYAN)
    text_box(slide, 0.55, 7.15, 11.5, 0.22,
             "Chest X-ray disease detection  •  fixed test split  •  measured claims only",
             size=9, color=MUTED)
    text_box(slide, 12.16, 7.1, 0.5, 0.3, str(len(prs.slides)),
             size=10, color=MUTED, align=PP_ALIGN.RIGHT)
    return slide


def card(slide, x, y, w, h, heading, detail, *, accent=CYAN):
    rect(slide, x, y, w, h, PANEL, radius=True)
    rect(slide, x, y, 0.075, h, accent)
    text_box(slide, x + 0.23, y + 0.18, w - 0.46, 0.42,
             heading, size=20, color=accent, bold=True)
    text_box(slide, x + 0.23, y + 0.72, w - 0.46, h - 0.88,
             detail, size=15)


def row(slide, y, columns, widths, *, header=False, color=WHITE):
    x = 0.67
    rect(slide, 0.55, y, 12.2, 0.55, PANEL if header else NAVY)
    for value, width in zip(columns, widths):
        text_box(slide, x, y + 0.05, width - 0.13, 0.42, str(value),
                 size=14 if header else 13, color=CYAN if header else color,
                 bold=header)
        x += width


def build() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    summary = pd.read_csv(ROOT / "artifacts/cnn_closeout/cnn_multiseed_summary.csv")
    efficiency = pd.read_csv(ROOT / "artifacts/T35_efficiency/architecture_only_efficiency.csv")
    paired = pd.read_csv(ROOT / "artifacts/explainable_ai/paired_eil_bootstrap.csv")

    s = base(prs, "Can lung-guided models look in the right place?", "TRUSTWORTHY CXR CLASSIFICATION")
    text_box(s, 0.65, 2.1, 11.7, 1.0,
             "Classification  •  anatomical localisation  •  actual input dependence",
             size=25, color=CYAN, bold=True)
    text_box(s, 0.65, 3.38, 10.9, 1.0,
             "Three-CNN closeout is measured. The matched perturbation experiment is prepared but not yet run.",
             size=23)
    text_box(s, 0.65, 5.56, 11.6, 0.7,
             "Ravindu-Pathirana / Chest-X-ray-Disease-Detection • explainable-AI branch",
             size=15, color=MUTED)

    s = base(prs, "The evidence ladder")
    card(s, 0.6, 2.0, 2.9, 3.7, "L1  Alignment", "Does the learned attention map overlap the lung mask?")
    card(s, 3.66, 2.0, 2.9, 3.7, "L2  Localisation", "Does predicted-class Grad-CAM energy move into lungs?")
    card(s, 6.72, 2.0, 2.9, 3.7, "L3  Dependence", "Does changing lung versus background content change predictions?", accent=CORAL)
    card(s, 9.78, 2.0, 2.9, 3.7, "L4  Transfer", "Does performance hold on a separately mapped external dataset?")
    text_box(s, 0.7, 6.15, 11.8, 0.5,
             "Higher EIL is evidence about heatmap location—not proof of lower background reliance.",
             size=18, color=CORAL, bold=True)

    s = base(prs, "Fixed data and comparison design")
    card(s, 0.7, 2.0, 3.75, 2.75, "21,165 images", "Four CXR classes; image-level split because patient identifiers are unavailable.")
    card(s, 4.8, 2.0, 3.75, 2.75, "14,815 / 3,175 / 3,175",
         "Train / validation / test. One committed manifest reused across experiments.")
    card(s, 8.9, 2.0, 3.75, 2.75, "1,000 fixed CAM images",
         "The same subset and image IDs support paired CNN EIL comparisons.")
    text_box(s, 0.75, 5.45, 11.7, 0.85,
             "Seed-42 ablations + three matched seeds for the selected CNN pairs.", size=21)

    s = base(prs, "Three candidates, one transparent selection boundary")
    card(s, 0.65, 2.0, 3.9, 3.62, "A  Residual lung gate",
         "Soft-mask-supervised spatial gate after backbone features. No mask needed at inference.")
    card(s, 4.72, 2.0, 3.9, 3.62, "B  Auxiliary segmentation",
         "Full-resolution lung-mask decoder with a classification + BCE/Dice objective.")
    card(s, 8.79, 2.0, 3.9, 3.62, "C  Activation penalty",
         "Penalises ReLU feature activation outside lungs during training; not a Grad-CAM loss.", accent=CORAL)
    text_box(s, 0.7, 6.05, 12, 0.58,
             "C led the DenseNet classification ablation; A was carried forward for its measured EIL and lightweight transferable gate—not declared a global winner.",
             size=16, color=MUTED)

    s = base(prs, "Three-CNN results: matched-seed means")
    row(s, 1.94, ["Backbone", "A0 → selected", "Accuracy Δ", "Macro-F1 Δ", "EIL Δ"],
        [2.25, 3.25, 2.05, 2.2, 2.45], header=True)
    for i, r in enumerate(summary.itertuples(index=False)):
        row(s, 2.65 + i * 0.77,
            [r.backbone, f"{r.baseline_arm} → {r.selected_arm}",
             f"{100*r.delta_accuracy_mean:+.2f} pp",
             f"{100*r.delta_macro_f1_mean:+.2f} pp",
             f"{100*r.delta_eil_mean:+.2f} pp"],
            [2.25, 3.25, 2.05, 2.2, 2.45])
    text_box(s, 0.72, 5.48, 11.8, 0.9,
             "EIL rose for all three selected arms, while classification changes were mixed. These are anatomical-localisation findings.",
             size=20, color=CORAL, bold=True)
    text_box(s, 0.7, 6.64, 11.6, 0.35,
             "Source: artifacts/cnn_closeout/cnn_multiseed_summary.csv", size=10, color=MUTED)

    s = base(prs, "Paired image-level EIL gains")
    seed42 = paired.loc[paired["seed"] == 42]
    for i, r in enumerate(seed42.itertuples(index=False)):
        label = {"DenseNet121": "DenseNet121", "ResNet50": "ResNet50",
                 "EfficientNet-B0": "EfficientNet-B0"}.get(r.backbone, r.backbone)
        y = 2.0 + i * 1.45
        rect(s, 0.72, y, 11.8, 1.16, PANEL, radius=True)
        rect(s, 0.72, y, 0.075, 1.16, CYAN)
        text_box(s, 0.98, y + 0.18, 3.0, 0.7, label, size=21, color=CYAN, bold=True)
        text_box(s, 4.0, y + 0.16, 8.2, 0.72,
                 f"Δ EIL {r.delta_eil_mean:+.3f}  |  95% CI [{r.ci_95_low:+.3f}, {r.ci_95_high:+.3f}]",
                 size=20)
    text_box(s, 0.74, 6.62, 11.5, 0.3,
             "1,000 shared images; 2,000 paired resamples. Retrospective precision analysis, not preregistered.",
             size=11, color=MUTED)

    s = base(prs, "The missing test: matched input dependence")
    card(s, 0.7, 2.0, 5.75, 3.55, "Primary interventions",
         "Lung blur ↔ background blur; lung swap ↔ background swap. A fixed different-class donor and equal-pixel-count background control are reused across arms.")
    card(s, 6.7, 2.0, 5.75, 3.55, "Primary measure",
         "TV distance dP = ½ Σ |p(original) − p(perturbed)|. LRG = dP(lung) − dP(background), paired by image and mode.", accent=CORAL)
    text_box(s, 0.75, 6.08, 11.7, 0.58,
             "Implementation and 3,175 donor pairs are ready; no new CXR inference has been run.",
             size=18, color=MUTED)

    s = base(prs, "T35 efficiency: what is measured now")
    row(s, 1.92, ["Backbone", "Params A0 → module", "GFLOPs* A0 → module", "CPU ms A0 → module"],
        [2.25, 3.18, 3.18, 3.59], header=True)
    for i, backbone in enumerate(("densenet121", "resnet50", "efficientnet_b0", "vit_base_patch16_224")):
        pair = efficiency[efficiency["backbone"] == backbone]
        b = pair.iloc[0]
        m = pair.iloc[1]
        label = {"densenet121": "DenseNet121", "resnet50": "ResNet50",
                 "efficientnet_b0": "EfficientNet-B0",
                 "vit_base_patch16_224": "ViT-Base/16"}[backbone]
        row(s, 2.53 + i * 0.64,
            [label, f"{b.params_total/1e6:.2f}M → {m.params_total/1e6:.2f}M",
             f"{b.gflops:.2f} → {m.gflops:.2f}",
             f"{b.cpu_latency_ms:.0f} → {m.cpu_latency_ms:.0f}"],
            [2.25, 3.18, 3.18, 3.59])
    text_box(s, 0.72, 5.55, 11.9, 0.9,
             "Same-session CPU, random-weight architecture benchmark. *FLOPs are partial fvcore estimates; fused attention and other ops were unsupported. Final GPU latency and trained-checkpoint size remain pending.",
             size=16, color=CORAL)
    text_box(s, 0.72, 6.62, 11.5, 0.32,
             "Source: artifacts/T35_efficiency/architecture_only_efficiency.csv", size=10, color=MUTED)

    s = base(prs, "What remains before a full trustworthiness claim")
    card(s, 0.7, 2.0, 5.75, 3.7, "Inputs and verification",
         "Primary images + masks; ResNet A0/A3 and ViT weights; secure Kaggle access. Each checkpoint needs strict load, 100-image smoke and reproduced test accuracy.")
    card(s, 6.7, 2.0, 5.75, 3.7, "Measured work still pending",
         "CNN L3 dependence; ViT explanation validation; validation-fitted calibration; RSNA mapping/inference; controlled all-model efficiency; final statistics and figures.", accent=CORAL)
    text_box(s, 0.75, 6.2, 11.5, 0.48,
             "Missing results remain NOT MEASURED—not zero and not inferred from EIL.",
             size=18, color=CORAL, bold=True)

    s = base(prs, "Takeaway")
    text_box(s, 0.78, 2.18, 11.6, 1.0,
             "Lung guidance consistently shifted CNN Grad-CAM energy toward lungs.",
             size=29, color=CYAN, bold=True)
    text_box(s, 0.78, 3.56, 11.6, 1.3,
             "Whether that shift reduces reliance on background shortcuts is still an empirical question.",
             size=28, color=CORAL, bold=True)
    text_box(s, 0.8, 5.67, 11.4, 0.7,
             "Next evidence: matched blur/swap perturbations, paired CIs, then external and efficiency checks.",
             size=19)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"Wrote {len(prs.slides)} slides to {OUT}")


if __name__ == "__main__":
    build()
